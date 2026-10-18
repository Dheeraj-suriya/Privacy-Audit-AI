"""Document text extraction with OCR fallback."""
import io
import re
from pathlib import Path
from typing import Tuple, Optional

from ..core.logging import logger
from ..database.models import Document

class ExtractionService:
    """Extract text from various document formats."""

    def extract(self, document: Document, file_path: Path) -> Tuple[str, str, Optional[float]]:
        """
        Extract text from document.
        Returns: (extracted_text, method, ocr_confidence)
        """
        file_type = document.file_type

        try:
            if file_type == "pdf":
                return self._extract_pdf(file_path)
            elif file_type == "docx":
                return self._extract_docx(file_path)
            elif file_type == "xlsx":
                return self._extract_xlsx(file_path)
            elif file_type == "csv":
                return self._extract_csv(file_path)
            elif file_type == "txt":
                return self._extract_txt(file_path)
            elif file_type == "image":
                return self._extract_image(file_path)
            else:
                return "", "unsupported", None
        except Exception as e:
            logger.error(f"Extraction failed for {document.id}: {e}")
            return "", "failed", None

    def _extract_pdf(self, file_path: Path) -> Tuple[str, str, Optional[float]]:
        """Extract text from PDF, fallback to OCR."""
        text = ""
        method = "pdfplumber"
        confidence = None

        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    text += page_text + "\n"

            # If text is too short, try OCR
            if len(text.strip()) < 100:
                logger.info(f"PDF text too short ({len(text)} chars), triggering OCR")
                ocr_text, confidence = self._ocr_pdf(file_path)
                if len(ocr_text) > len(text):
                    text = ocr_text
                    method = "ocr"
        except ImportError:
            logger.warning("pdfplumber not available, trying OCR")
            text, confidence = self._ocr_pdf(file_path)
            method = "ocr"
        except Exception as e:
            logger.error(f"PDF extraction error: {e}")
            text, confidence = self._ocr_pdf(file_path)
            method = "ocr"

        return text, method, confidence

    def _ocr_pdf(self, file_path: Path) -> Tuple[str, Optional[float]]:
        """OCR extraction for PDFs via image conversion."""
        try:
            from pdf2image import convert_from_path
            import pytesseract

            images = convert_from_path(str(file_path), dpi=200)
            text_parts = []
            total_confidence = 0.0

            for img in images:
                # Get OCR data with confidence
                data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
                confidences = [int(c) for c in data["conf"] if int(c) > 0]
                if confidences:
                    total_confidence += sum(confidences) / len(confidences)

                page_text = pytesseract.image_to_string(img)
                text_parts.append(page_text)

            full_text = "\n".join(text_parts)
            avg_confidence = total_confidence / len(images) if images else None

            return full_text, avg_confidence
        except Exception as e:
            logger.error(f"OCR failed: {e}")
            return "", None

    def _extract_docx(self, file_path: Path) -> Tuple[str, str, Optional[float]]:
        """Extract text from DOCX including tables."""
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)

            text_parts = []

            # Paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_parts.append(row_text)

            return "\n".join(text_parts), "python-docx", None
        except ImportError:
            logger.warning("python-docx not available")
            return "", "unavailable", None
        except Exception as e:
            logger.error(f"DOCX extraction error: {e}")
            return "", "failed", None

    def _extract_xlsx(self, file_path: Path) -> Tuple[str, str, Optional[float]]:
        """Extract text from XLSX."""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text_parts = []

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                text_parts.append(f"Sheet: {sheet_name}")

                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(cell) for cell in row if cell is not None)
                    if row_text:
                        text_parts.append(row_text)

            return "\n".join(text_parts), "openpyxl", None
        except ImportError:
            logger.warning("openpyxl not available")
            return "", "unavailable", None
        except Exception as e:
            logger.error(f"XLSX extraction error: {e}")
            return "", "failed", None

    def _extract_csv(self, file_path: Path) -> Tuple[str, str, Optional[float]]:
        """Extract text from CSV."""
        try:
            import csv
            text_parts = []
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                for row in reader:
                    text_parts.append(" | ".join(row))
            return "\n".join(text_parts), "csv", None
        except Exception as e:
            logger.error(f"CSV extraction error: {e}")
            return "", "failed", None

    def _extract_txt(self, file_path: Path) -> Tuple[str, str, Optional[float]]:
        """Extract text from plain text file."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read(), "text", None
        except Exception as e:
            logger.error(f"Text extraction error: {e}")
            return "", "failed", None

    def _extract_image(self, file_path: Path) -> Tuple[str, str, Optional[float]]:
        """Extract text from image using Tesseract OCR."""
        try:
            from PIL import Image
            import pytesseract

            img = Image.open(file_path)

            # Get confidence data
            data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
            confidences = [int(c) for c in data["conf"] if int(c) > 0]
            avg_confidence = sum(confidences) / len(confidences) if confidences else None

            text = pytesseract.image_to_string(img)
            return text, "tesseract-ocr", avg_confidence
        except ImportError:
            logger.warning("Tesseract OCR not available")
            return "", "unavailable", None
        except Exception as e:
            logger.error(f"Image OCR error: {e}")
            return "", "failed", None

extraction_service = ExtractionService()
